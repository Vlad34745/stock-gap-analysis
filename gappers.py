"""
Stock Gap & Price Action Analysis Tool
----------------------------------------
Fetches historical market data, detects opening gaps above a threshold,
and exports a formatted Excel report with subsequent price action analysis,
gap-fill detection, and an embedded chart.

By default only gap-ups (Open > previous Close) are detected. Use
--direction to also include gap-downs or both.

Multiple tickers can be analyzed in one run (comma-separated); each
gets its own sheet in a single output workbook, plus a Summary
comparison sheet.

Usage:
    python gappers.py --ticker TSLA --start 2026-01-01 --end 2026-05-15
    python gappers.py --ticker AAPL --threshold 0.02
    python gappers.py --ticker AAPL --direction both
    python gappers.py --ticker AAPL,TSLA,NVDA --output multi_report.xlsx
    python gappers.py --ticker AAPL --csv
"""

import argparse
import hashlib
import logging
import time
from pathlib import Path
from typing import Optional

import yfinance as yf
import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (
    Paragraph, ParagraphProperties, CharacterProperties, Font as DrawingFont, RichTextProperties
)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)
logger = logging.getLogger(__name__)

CACHE_DIR = Path(".cache")
CACHE_MAX_AGE_SECONDS = 24 * 60 * 60  # 1 day
MAX_FETCH_RETRIES = 3
RETRY_BACKOFF_SECONDS = 2


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Stock gap and price action analyzer.")
    parser.add_argument(
        "--ticker", default="AAPL",
        help="Ticker symbol, e.g. TSLA, NVDA, BTC-USD. Comma-separated for multiple, e.g. AAPL,TSLA,NVDA"
    )
    parser.add_argument("--start", default="2026-01-01", help="Start date (YYYY-MM-DD)")
    parser.add_argument("--end", default="2026-05-15", help="End date (YYYY-MM-DD)")
    parser.add_argument(
        "--threshold", type=float, default=0.015,
        help="Minimum opening gap to include, as a decimal (0.015 = 1.5%%)"
    )
    parser.add_argument(
        "--direction", choices=["up", "down", "both"], default="up",
        help="Which gap direction to detect: 'up' (default), 'down', or 'both'"
    )
    parser.add_argument("--output", default=None, help="Output .xlsx filename")
    parser.add_argument(
        "--no-cache", action="store_true",
        help="Skip local cache and always re-fetch fresh data from Yahoo Finance"
    )
    parser.add_argument(
        "--csv", action="store_true",
        help="Also save the raw gap events (all tickers combined) to a .csv file next to the .xlsx report"
    )
    args = parser.parse_args()

    if args.threshold <= 0:
        parser.error("--threshold must be a positive number (e.g. 0.015 for 1.5%)")

    try:
        start_date = pd.to_datetime(args.start)
        end_date = pd.to_datetime(args.end)
    except ValueError as exc:
        parser.error(f"Invalid date format, expected YYYY-MM-DD: {exc}")
    else:
        if start_date >= end_date:
            parser.error(f"--start ({args.start}) must be before --end ({args.end})")

    return args


def _cache_path(ticker: str, start: str, end: str) -> Path:
    key = hashlib.md5(f"{ticker}_{start}_{end}".encode()).hexdigest()[:12]
    return CACHE_DIR / f"{ticker}_{key}.csv"


def _load_from_cache(path: Path) -> Optional[pd.DataFrame]:
    if not path.exists():
        return None
    age = time.time() - path.stat().st_mtime
    if age > CACHE_MAX_AGE_SECONDS:
        return None
    try:
        data = pd.read_csv(path, index_col=0, parse_dates=True)
    except Exception:
        return None
    return data if not data.empty else None


def fetch_data(ticker: str, start: str, end: str, use_cache: bool = True) -> Optional[pd.DataFrame]:
    cache_path = _cache_path(ticker, start, end)

    if use_cache:
        cached = _load_from_cache(cache_path)
        if cached is not None:
            logger.info("Using cached data for %s (from %s).", ticker, cache_path.name)
            return cached

    data = None
    for attempt in range(1, MAX_FETCH_RETRIES + 1):
        logger.info(
            "Fetching historical data for %s from Yahoo Finance (attempt %d/%d)...",
            ticker, attempt, MAX_FETCH_RETRIES
        )
        try:
            data = yf.download(ticker, start=start, end=end, auto_adjust=False, actions=False)
        except Exception as exc:
            logger.warning("Attempt %d failed for '%s': %s", attempt, ticker, exc)
            data = None

        if data is not None and isinstance(data.columns, pd.MultiIndex):
            data.columns = data.columns.droplevel(1)

        if data is not None and not data.empty:
            break

        if attempt < MAX_FETCH_RETRIES:
            wait = RETRY_BACKOFF_SECONDS * attempt
            logger.info("Retrying in %ds...", wait)
            time.sleep(wait)

    if data is None or data.empty:
        logger.error("No data returned for ticker '%s' after %d attempts. Check the symbol or date range.",
                     ticker, MAX_FETCH_RETRIES)
        return None

    if use_cache:
        try:
            CACHE_DIR.mkdir(exist_ok=True)
            data.to_csv(cache_path)
        except Exception as exc:
            logger.warning("Could not write cache for '%s': %s", ticker, exc)

    return data


def calculate_gaps(data: pd.DataFrame, threshold: float, direction: str = "up") -> pd.DataFrame:
    data = data.copy()
    data["Prev_Close"] = data["Close"].shift(1)
    data["Gap_Pct"] = (data["Open"] - data["Prev_Close"]) / data["Prev_Close"]
    data["Day2_Move_Pct"] = (data["Close"].shift(-1) - data["Close"]) / data["Close"]
    data["Day3_Move_Pct"] = (data["Close"].shift(-2) - data["Close"]) / data["Close"]

    # A gap "fills" when price trades back to the previous close on the same day:
    # for a gap-up, that means the day's Low dipped back down to Prev_Close;
    # for a gap-down, it means the day's High climbed back up to Prev_Close.
    gap_up_filled = data["Low"] <= data["Prev_Close"]
    gap_down_filled = data["High"] >= data["Prev_Close"]
    data["Gap_Filled"] = np.where(data["Gap_Pct"] > 0, gap_up_filled, gap_down_filled)

    if direction == "up":
        mask = data["Gap_Pct"] > threshold
    elif direction == "down":
        mask = data["Gap_Pct"] < -threshold
    else:  # both
        mask = data["Gap_Pct"].abs() > threshold

    gappers = data[mask].dropna(subset=["Prev_Close"])
    logger.info(
        "Found %d gap events (direction=%s) above %.2f%% threshold.",
        len(gappers), direction, threshold * 100
    )
    return gappers


ACCENT_TEAL = "004D40"
ACCENT_GREEN = "2E7D32"
ACCENT_RED = "C62828"
GRID_GRAY = "E0E0E0"


def _styled_title(text: str, size: int = 1400, color: str = ACCENT_TEAL):
    """Build a RichText chart title in the report's font/color instead of Excel's plain default."""
    cp = CharacterProperties(sz=size, b=True, solidFill=color, latin=DrawingFont(typeface="Segoe UI"))
    pp = ParagraphProperties(defRPr=cp)
    return _title_from_text(text, cp, pp)


def _title_from_text(text, cp, pp):
    from openpyxl.chart.title import Title
    from openpyxl.chart.text import Text
    from openpyxl.drawing.text import RegularTextRun
    run = RegularTextRun(rPr=cp, t=text)
    para = Paragraph(pPr=pp, r=[run])
    rich = RichText(bodyPr=RichTextProperties(), p=[para])
    return Title(tx=Text(rich=rich))


def _style_axis_text(axis, size: int = 900, color: str = "424242", rot: int = -2700000):
    cp = CharacterProperties(sz=size, solidFill=color, latin=DrawingFont(typeface="Segoe UI"))
    pp = ParagraphProperties(defRPr=cp)
    axis.txPr = RichText(bodyPr=RichTextProperties(rot=rot, vert="horz"), p=[Paragraph(pPr=pp, r=[])])


def _style_axis_title(axis, text: str, size: int = 1000, color: str = ACCENT_TEAL):
    cp = CharacterProperties(sz=size, b=True, solidFill=color, latin=DrawingFont(typeface="Segoe UI"))
    pp = ParagraphProperties(defRPr=cp)
    axis.title = _title_from_text(text, cp, pp)


def _light_gridlines():
    return ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill=GRID_GRAY, w=6350)))


def build_report(gappers: pd.DataFrame, ticker: str, wb: openpyxl.Workbook = None) -> openpyxl.Workbook:
    """
    Build (or append to) a formatted gap-analysis report: a color-coded
    table (including gap-fill detection and a Fill Rate figure) plus an
    embedded bar chart of Gap % across all detected events.

    If `wb` is None, a new workbook is created and the sheet replaces the
    default blank sheet. If `wb` is given, a new sheet is appended to it —
    this is how multiple tickers end up in one file, one sheet each.
    """
    if wb is None:
        wb = openpyxl.Workbook()
        ws = wb.active
    else:
        ws = wb.create_sheet()

    # Excel sheet names: max 31 chars, no \ / ? * [ ] :
    safe_title = "".join(c for c in ticker if c not in r'\/?*[]:')[:31]
    ws.title = safe_title or "Sheet"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Segoe UI", size=16, bold=True, color="004D40")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=11, bold=True)

    fill_header = PatternFill(start_color="004D40", end_color="004D40", fill_type="solid")
    fill_zebra = PatternFill(start_color="F0F7F4", end_color="F0F7F4", fill_type="solid")
    fill_green = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    fill_red = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"),
    )

    ws["A1"] = f"Historical Stock Gap Analysis: {ticker}"
    ws["A1"].font = font_title
    ws.merge_cells("A1:H1")
    ws.row_dimensions[1].height = 30

    headers = ["Date", "Open Price", "Prev Close", "Gap %", "Filled?", "Day 1 Close", "Day 2 Move", "Day 3 Move"]
    for idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 25

    current_row = 5
    for date, row in gappers.iterrows():
        is_zebra = current_row % 2 == 0
        row_fill = fill_zebra if is_zebra else PatternFill(fill_type=None)

        ws.cell(row=current_row, column=1, value=date.strftime("%Y-%m-%d")).alignment = \
            Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=2, value=float(row["Open"])).number_format = "$#,##0.00"
        ws.cell(row=current_row, column=3, value=float(row["Prev_Close"])).number_format = "$#,##0.00"

        c_gap = ws.cell(row=current_row, column=4, value=float(row["Gap_Pct"]))
        c_gap.number_format = "0.00%"
        if row["Gap_Pct"] > 0:
            c_gap.font = Font(name="Segoe UI", size=11, bold=True, color="1B5E20")
            c_gap.fill = fill_green
        else:
            c_gap.font = Font(name="Segoe UI", size=11, bold=True, color="B71C1C")
            c_gap.fill = fill_red

        c_filled = ws.cell(row=current_row, column=5, value="Yes" if bool(row["Gap_Filled"]) else "No")
        c_filled.alignment = Alignment(horizontal="center", vertical="center")
        if bool(row["Gap_Filled"]):
            c_filled.font = Font(name="Segoe UI", size=11, color="1B5E20", bold=True)
            c_filled.fill = fill_green
        else:
            c_filled.font = Font(name="Segoe UI", size=11, color="757575")

        ws.cell(row=current_row, column=6, value=float(row["Close"])).number_format = "$#,##0.00"

        for col_idx, col_name in [(7, "Day2_Move_Pct"), (8, "Day3_Move_Pct")]:
            raw_val = row[col_name]
            c_move = ws.cell(row=current_row, column=col_idx)
            if pd.isna(raw_val):
                c_move.value = "N/A"
                c_move.font = Font(name="Segoe UI", size=11, color="9E9E9E", italic=True)
            else:
                val = float(raw_val)
                c_move.value = val
                c_move.number_format = "0.00%"
                if val > 0:
                    c_move.fill = fill_green
                    c_move.font = Font(name="Segoe UI", size=11, color="1B5E20", bold=True)
                else:
                    c_move.fill = fill_red
                    c_move.font = Font(name="Segoe UI", size=11, color="B71C1C")

        for col in range(1, 9):
            cell = ws.cell(row=current_row, column=col)
            cell.border = thin_border
            if col != 1:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            if col not in (4, 5, 7, 8) and is_zebra:
                cell.fill = row_fill

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    summary_row = current_row + 1
    ws.cell(row=summary_row, column=1, value="Average").font = font_bold
    ws.cell(row=summary_row, column=1).alignment = Alignment(horizontal="left", vertical="center")

    for col_idx, letter in [(4, "D"), (7, "G"), (8, "H")]:
        c_avg = ws.cell(row=summary_row, column=col_idx, value=f"=AVERAGE({letter}5:{letter}{summary_row - 2})")
        c_avg.font = font_bold
        c_avg.number_format = "0.00%"
        c_avg.alignment = Alignment(horizontal="right", vertical="center")
        c_avg.border = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

    fill_rate = float(gappers["Gap_Filled"].mean()) if len(gappers) else 0.0
    c_fill_rate = ws.cell(row=summary_row, column=5, value=fill_rate)
    c_fill_rate.font = font_bold
    c_fill_rate.number_format = "0.0%"
    c_fill_rate.alignment = Alignment(horizontal="center", vertical="center")
    c_fill_rate.border = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

    if len(gappers) >= 1:
        chart = BarChart()
        chart.type = "col"
        chart.gapWidth = 40
        chart.title = _styled_title(f"{ticker} — Gap % by Event")
        chart.style = None
        chart.height = 9
        chart.width = 19
        chart.legend = None

        data_ref = Reference(ws, min_col=4, min_row=4, max_row=current_row - 1)
        cats_ref = Reference(ws, min_col=1, min_row=5, max_row=current_row - 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        # Color each bar green (gap up) or red (gap down) instead of Excel's
        # default single flat color, so direction is visible at a glance.
        series = chart.series[0]
        series.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
        series.data_points = [
            DataPoint(idx=i, spPr=GraphicalProperties(
                solidFill=ACCENT_GREEN if val > 0 else ACCENT_RED,
                ln=LineProperties(noFill=True),
            ))
            for i, val in enumerate(gappers["Gap_Pct"].tolist())
        ]

        chart.y_axis.number_format = "0.0%"
        chart.y_axis.majorGridlines = _light_gridlines()
        chart.y_axis.delete = False
        _style_axis_title(chart.y_axis, "Gap %")
        _style_axis_text(chart.y_axis, size=900)

        chart.x_axis.delete = False
        chart.x_axis.majorGridlines = None
        chart.x_axis.tickLblPos = "low"
        _style_axis_title(chart.x_axis, "Date")
        _style_axis_text(chart.x_axis, size=800)

        chart.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill=GRID_GRAY, w=6350))

        ws.add_chart(chart, "J4")

    return wb


def build_summary_sheet(wb: openpyxl.Workbook, stats: list) -> None:
    """
    Insert a 'Summary' sheet as the first sheet, comparing all analyzed
    tickers side by side: gap count, average Gap %, fill rate, and
    average Day2/Day3 follow-through — plus a bar chart comparing
    average Gap % across tickers.
    `stats` is a list of dicts, one per ticker.
    """
    ws = wb.create_sheet("Summary", 0)
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Segoe UI", size=16, bold=True, color="004D40")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="004D40", end_color="004D40", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"),
    )

    ws["A1"] = "Gap Analysis — Ticker Comparison"
    ws["A1"].font = font_title
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 30

    headers = ["Ticker", "Gap Events", "Avg Gap %", "Fill Rate", "Avg Day 2 Move", "Avg Day 3 Move"]
    for idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for i, s in enumerate(stats, start=4):
        ws.cell(row=i, column=1, value=s["ticker"]).font = Font(name="Segoe UI", bold=True)
        ws.cell(row=i, column=2, value=s["count"])
        for col, key, fmt in [
            (3, "avg_gap", "0.00%"), (4, "fill_rate", "0.0%"),
            (5, "avg_day2", "0.00%"), (6, "avg_day3", "0.00%"),
        ]:
            c = ws.cell(row=i, column=col, value=s[key])
            c.number_format = fmt
        for col in range(1, 7):
            ws.cell(row=i, column=col).border = thin_border
            if col != 1:
                ws.cell(row=i, column=col).alignment = Alignment(horizontal="right")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 15)

    if stats:
        chart = BarChart()
        chart.type = "col"
        chart.gapWidth = 60
        chart.title = _styled_title("Avg Gap % by Ticker")
        chart.style = None
        chart.height = 9
        chart.width = 19
        chart.legend = None

        last_row = 3 + len(stats)
        data_ref = Reference(ws, min_col=3, min_row=3, max_row=last_row)
        cats_ref = Reference(ws, min_col=1, min_row=4, max_row=last_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        series = chart.series[0]
        series.graphicalProperties = GraphicalProperties(
            solidFill=ACCENT_TEAL, ln=LineProperties(noFill=True)
        )

        chart.y_axis.number_format = "0.0%"
        chart.y_axis.majorGridlines = _light_gridlines()
        chart.y_axis.delete = False
        _style_axis_title(chart.y_axis, "Avg Gap %")
        _style_axis_text(chart.y_axis, size=900)

        chart.x_axis.delete = False
        chart.x_axis.majorGridlines = None
        chart.x_axis.tickLblPos = "low"
        _style_axis_title(chart.x_axis, "Ticker")
        _style_axis_text(chart.x_axis, size=900, rot=0)

        chart.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill=GRID_GRAY, w=6350))

        ws.add_chart(chart, "H3")


def main() -> None:
    args = parse_args()
    tickers = [t.strip().upper() for t in args.ticker.split(",") if t.strip()]

    wb = None
    tickers_with_data = []
    stats = []
    csv_frames = []
    for ticker in tickers:
        data = fetch_data(ticker, args.start, args.end, use_cache=not args.no_cache)
        if data is None:
            continue
        gappers = calculate_gaps(data, args.threshold, args.direction)

        if gappers.empty:
            logger.warning("No gap events found for '%s'. Skipping.", ticker)
            continue

        wb = build_report(gappers, ticker, wb)
        tickers_with_data.append(ticker)

        fill_rate = float(gappers["Gap_Filled"].mean())
        stats.append({
            "ticker": ticker,
            "count": len(gappers),
            "avg_gap": float(gappers["Gap_Pct"].mean()),
            "fill_rate": fill_rate,
            "avg_day2": float(gappers["Day2_Move_Pct"].dropna().mean()) if gappers["Day2_Move_Pct"].notna().any() else 0.0,
            "avg_day3": float(gappers["Day3_Move_Pct"].dropna().mean()) if gappers["Day3_Move_Pct"].notna().any() else 0.0,
        })

        if args.csv:
            labeled = gappers.copy()
            labeled.insert(0, "Ticker", ticker)
            csv_frames.append(labeled)

    if wb is None:
        logger.warning("No gap events found for any ticker. No report generated.")
        return

    if len(tickers_with_data) > 1:
        build_summary_sheet(wb, stats)

    if args.output:
        output_filename = args.output
    elif len(tickers_with_data) == 1:
        output_filename = f"{tickers_with_data[0]}_Gap_Analysis_Project.xlsx"
    else:
        output_filename = "Multi_Ticker_Gap_Analysis_Project.xlsx"

    wb.save(output_filename)
    logger.info("Success! Report saved as: %s (%d ticker sheet(s))", output_filename, len(tickers_with_data))

    if args.csv and csv_frames:
        csv_filename = str(Path(output_filename).with_suffix(".csv"))
        pd.concat(csv_frames).to_csv(csv_filename)
        logger.info("Raw gap events also saved as: %s", csv_filename)

    print("\n=== Gap Analysis Summary ===")
    for s in stats:
        print(
            f"{s['ticker']:<8} {s['count']:>3} gaps | "
            f"avg gap {s['avg_gap']:+.2%} | fill rate {s['fill_rate']:.1%} | "
            f"avg Day2 {s['avg_day2']:+.2%} | avg Day3 {s['avg_day3']:+.2%}"
        )


if __name__ == "__main__":
    main()