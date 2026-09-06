"""Building the formatted Excel report: per-ticker sheets and the multi-ticker Summary sheet."""
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

from .config import ACCENT_TEAL, ACCENT_GREEN, ACCENT_RED, GRID_GRAY
from .chart_style import styled_title, style_axis_text, style_axis_title, light_gridlines


def build_report(gappers: pd.DataFrame, ticker: str, wb: openpyxl.Workbook = None) -> openpyxl.Workbook:
    """
    Build (or append to) a formatted gap-analysis report: a color-coded
    table (including gap-fill detection and a Fill Rate figure) plus an
    embedded bar chart of Gap % across all detected events.

    If `wb` is None, a new workbook is created and the sheet replaces the
    default blank sheet. If `wb` is given, a new sheet is appended to it —
    this is how multiple tickers end up in one file, one sheet each.
    """
    if wb is None:
        wb = openpyxl.Workbook()
        ws = wb.active
    else:
        ws = wb.create_sheet()

    # Excel sheet names: max 31 chars, no \ / ? * [ ] :
    safe_title = "".join(c for c in ticker if c not in r'\/?*[]:')[:31]
    ws.title = safe_title or "Sheet"
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Segoe UI", size=16, bold=True, color=ACCENT_TEAL)
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=11, bold=True)

    fill_header = PatternFill(start_color=ACCENT_TEAL, end_color=ACCENT_TEAL, fill_type="solid")
    fill_zebra = PatternFill(start_color="F0F7F4", end_color="F0F7F4", fill_type="solid")
    fill_green = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    fill_red = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"),
    )

    ws["A1"] = f"Historical Stock Gap Analysis: {ticker}"
    ws["A1"].font = font_title
    ws.merge_cells("A1:H1")
    ws.row_dimensions[1].height = 30

    headers = ["Date", "Open Price", "Prev Close", "Gap %", "Filled?", "Day 1 Close", "Day 2 Move", "Day 3 Move"]
    for idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 25

    current_row = 5
    for date, row in gappers.iterrows():
        is_zebra = current_row % 2 == 0
        row_fill = fill_zebra if is_zebra else PatternFill(fill_type=None)

        ws.cell(row=current_row, column=1, value=date.strftime("%Y-%m-%d")).alignment = \
            Alignment(horizontal="center", vertical="center")
        ws.cell(row=current_row, column=2, value=float(row["Open"])).number_format = "$#,##0.00"
        ws.cell(row=current_row, column=3, value=float(row["Prev_Close"])).number_format = "$#,##0.00"

        c_gap = ws.cell(row=current_row, column=4, value=float(row["Gap_Pct"]))
        c_gap.number_format = "0.00%"
        if row["Gap_Pct"] > 0:
            c_gap.font = Font(name="Segoe UI", size=11, bold=True, color="1B5E20")
            c_gap.fill = fill_green
        else:
            c_gap.font = Font(name="Segoe UI", size=11, bold=True, color="B71C1C")
            c_gap.fill = fill_red

        c_filled = ws.cell(row=current_row, column=5, value="Yes" if bool(row["Gap_Filled"]) else "No")
        c_filled.alignment = Alignment(horizontal="center", vertical="center")
        if bool(row["Gap_Filled"]):
            c_filled.font = Font(name="Segoe UI", size=11, color="1B5E20", bold=True)
            c_filled.fill = fill_green
        else:
            c_filled.font = Font(name="Segoe UI", size=11, color="757575")

        ws.cell(row=current_row, column=6, value=float(row["Close"])).number_format = "$#,##0.00"

        for col_idx, col_name in [(7, "Day2_Move_Pct"), (8, "Day3_Move_Pct")]:
            raw_val = row[col_name]
            c_move = ws.cell(row=current_row, column=col_idx)
            if pd.isna(raw_val):
                c_move.value = "N/A"
                c_move.font = Font(name="Segoe UI", size=11, color="9E9E9E", italic=True)
            else:
                val = float(raw_val)
                c_move.value = val
                c_move.number_format = "0.00%"
                if val > 0:
                    c_move.fill = fill_green
                    c_move.font = Font(name="Segoe UI", size=11, color="1B5E20", bold=True)
                else:
                    c_move.fill = fill_red
                    c_move.font = Font(name="Segoe UI", size=11, color="B71C1C")

        for col in range(1, 9):
            cell = ws.cell(row=current_row, column=col)
            cell.border = thin_border
            if col != 1:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            if col not in (4, 5, 7, 8) and is_zebra:
                cell.fill = row_fill

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    summary_row = current_row + 1
    ws.cell(row=summary_row, column=1, value="Average").font = font_bold
    ws.cell(row=summary_row, column=1).alignment = Alignment(horizontal="left", vertical="center")

    for col_idx, letter in [(4, "D"), (7, "G"), (8, "H")]:
        c_avg = ws.cell(row=summary_row, column=col_idx, value=f"=AVERAGE({letter}5:{letter}{summary_row - 2})")
        c_avg.font = font_bold
        c_avg.number_format = "0.00%"
        c_avg.alignment = Alignment(horizontal="right", vertical="center")
        c_avg.border = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

    fill_rate = float(gappers["Gap_Filled"].mean()) if len(gappers) else 0.0
    c_fill_rate = ws.cell(row=summary_row, column=5, value=fill_rate)
    c_fill_rate.font = font_bold
    c_fill_rate.number_format = "0.0%"
    c_fill_rate.alignment = Alignment(horizontal="center", vertical="center")
    c_fill_rate.border = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

    if len(gappers) >= 1:
        chart = BarChart()
        chart.type = "col"
        chart.gapWidth = 40
        chart.title = styled_title(f"{ticker} — Gap % by Event")
        chart.style = None
        chart.height = 9
        chart.width = 19
        chart.legend = None

        data_ref = Reference(ws, min_col=4, min_row=4, max_row=current_row - 1)
        cats_ref = Reference(ws, min_col=1, min_row=5, max_row=current_row - 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        # Color each bar green (gap up) or red (gap down) instead of Excel's
        # default single flat color, so direction is visible at a glance.
        series = chart.series[0]
        series.graphicalProperties = GraphicalProperties(ln=LineProperties(noFill=True))
        series.data_points = [
            DataPoint(idx=i, spPr=GraphicalProperties(
                solidFill=ACCENT_GREEN if val > 0 else ACCENT_RED,
                ln=LineProperties(noFill=True),
            ))
            for i, val in enumerate(gappers["Gap_Pct"].tolist())
        ]

        chart.y_axis.number_format = "0.0%"
        chart.y_axis.majorGridlines = light_gridlines()
        chart.y_axis.delete = False
        style_axis_title(chart.y_axis, "Gap %")
        style_axis_text(chart.y_axis, size=900)

        chart.x_axis.delete = False
        chart.x_axis.majorGridlines = None
        chart.x_axis.tickLblPos = "low"
        style_axis_title(chart.x_axis, "Date")
        style_axis_text(chart.x_axis, size=800)

        chart.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill=GRID_GRAY, w=6350))

        ws.add_chart(chart, "J4")

    return wb


def build_summary_sheet(wb: openpyxl.Workbook, stats: list) -> None:
    """
    Insert a 'Summary' sheet as the first sheet, comparing all analyzed
    tickers side by side: gap count, average Gap %, fill rate, and
    average Day2/Day3 follow-through — plus a bar chart comparing
    average Gap % across tickers.
    `stats` is a list of dicts, one per ticker.
    """
    ws = wb.create_sheet("Summary", 0)
    ws.views.sheetView[0].showGridLines = True

    font_title = Font(name="Segoe UI", size=16, bold=True, color=ACCENT_TEAL)
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color=ACCENT_TEAL, end_color=ACCENT_TEAL, fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CCCCCC"), right=Side(style="thin", color="CCCCCC"),
        top=Side(style="thin", color="CCCCCC"), bottom=Side(style="thin", color="CCCCCC"),
    )

    ws["A1"] = "Gap Analysis — Ticker Comparison"
    ws["A1"].font = font_title
    ws.merge_cells("A1:F1")
    ws.row_dimensions[1].height = 30

    headers = ["Ticker", "Gap Events", "Avg Gap %", "Fill Rate", "Avg Day 2 Move", "Avg Day 3 Move"]
    for idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for i, s in enumerate(stats, start=4):
        ws.cell(row=i, column=1, value=s["ticker"]).font = Font(name="Segoe UI", bold=True)
        ws.cell(row=i, column=2, value=s["count"])
        for col, key, fmt in [
            (3, "avg_gap", "0.00%"), (4, "fill_rate", "0.0%"),
            (5, "avg_day2", "0.00%"), (6, "avg_day3", "0.00%"),
        ]:
            c = ws.cell(row=i, column=col, value=s[key])
            c.number_format = fmt
        for col in range(1, 7):
            ws.cell(row=i, column=col).border = thin_border
            if col != 1:
                ws.cell(row=i, column=col).alignment = Alignment(horizontal="right")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 15)

    if stats:
        chart = BarChart()
        chart.type = "col"
        chart.gapWidth = 60
        chart.title = styled_title("Avg Gap % by Ticker")
        chart.style = None
        chart.height = 9
        chart.width = 19
        chart.legend = None

        last_row = 3 + len(stats)
        data_ref = Reference(ws, min_col=3, min_row=3, max_row=last_row)
        cats_ref = Reference(ws, min_col=1, min_row=4, max_row=last_row)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)

        series = chart.series[0]
        series.graphicalProperties = GraphicalProperties(
            solidFill=ACCENT_TEAL, ln=LineProperties(noFill=True)
        )

        chart.y_axis.number_format = "0.0%"
        chart.y_axis.majorGridlines = light_gridlines()
        chart.y_axis.delete = False
        style_axis_title(chart.y_axis, "Avg Gap %")
        style_axis_text(chart.y_axis, size=900)

        chart.x_axis.delete = False
        chart.x_axis.majorGridlines = None
        chart.x_axis.tickLblPos = "low"
        style_axis_title(chart.x_axis, "Ticker")
        style_axis_text(chart.x_axis, size=900, rot=0)

        chart.graphical_properties = GraphicalProperties(ln=LineProperties(solidFill=GRID_GRAY, w=6350))

        ws.add_chart(chart, "H3")
